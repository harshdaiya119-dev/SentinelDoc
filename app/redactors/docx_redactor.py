"""DOCX document redactor using python-docx."""
import io
from typing import Dict, List, Optional, Set, Tuple
import docx
from app.models.schemas import Finding
from .base import BaseRedactor


def _redact_in_paragraph(
    paragraph,
    target_text: str,
    start_char: Optional[int] = None,
    end_char: Optional[int] = None,
) -> bool:
    """Replace a specific occurrence of target_text with [REDACTED] in paragraph runs preserving styles."""
    if not target_text or not target_text.strip():
        return False

    p_text = paragraph.text
    if target_text not in p_text:
        return False

    # Determine exact target span [match_start, match_end]
    if (
        start_char is not None
        and end_char is not None
        and 0 <= start_char < end_char <= len(p_text)
        and p_text[start_char:end_char] == target_text
    ):
        match_start = start_char
        match_end = end_char
    elif start_char is not None:
        # Find all occurrences of target_text in p_text and pick the one closest to start_char
        positions = []
        idx = p_text.find(target_text)
        while idx != -1:
            positions.append(idx)
            idx = p_text.find(target_text, idx + 1)
        if positions:
            match_start = min(positions, key=lambda p: abs(p - start_char))
            match_end = match_start + len(target_text)
        else:
            return False
    else:
        match_start = p_text.find(target_text)
        if match_start == -1:
            return False
        match_end = match_start + len(target_text)

    # If no runs, fallback to single occurrence replacement in paragraph.text
    if not paragraph.runs:
        paragraph.text = p_text[:match_start] + "[REDACTED]" + p_text[match_end:]
        return True

    # Map each run to its [curr_pos, curr_pos + len] span in full paragraph text
    run_spans = []
    curr_pos = 0
    for run in paragraph.runs:
        r_len = len(run.text)
        run_spans.append((curr_pos, curr_pos + r_len, run))
        curr_pos += r_len

    # 1. Check if target_text span is fully within a single run
    for r_start, r_end, run in run_spans:
        if r_start <= match_start and match_end <= r_end:
            rel_start = match_start - r_start
            rel_end = match_end - r_start
            run.text = run.text[:rel_start] + "[REDACTED]" + run.text[rel_end:]
            return True

    # 2. Target text spans multiple runs
    first_overlapping = True
    for r_start, r_end, run in run_spans:
        if r_end <= match_start or r_start >= match_end:
            continue  # No overlap

        r_text = run.text
        rel_start = max(0, match_start - r_start)
        rel_end = min(len(r_text), match_end - r_start)

        before = r_text[:rel_start]
        after = r_text[rel_end:]

        if first_overlapping:
            run.text = before + "[REDACTED]"
            first_overlapping = False
        else:
            run.text = ""

        if match_end <= r_end:
            run.text += after

    return True


class DocxRedactor(BaseRedactor):
    """Redacts DOCX documents replacing PII with [REDACTED] in paragraphs and tables."""

    def redact(
        self,
        content: bytes,
        file_name: str,
        findings: List[Finding],
        target_finding_ids: Optional[Set[str]] = None,
    ) -> bytes:
        """Replace sensitive findings in DOCX paragraphs and tables with [REDACTED]."""
        if target_finding_ids is not None:
            to_redact = [f for f in findings if f.id in target_finding_ids]
        else:
            to_redact = list(findings)

        try:
            doc = docx.Document(io.BytesIO(content))
        except Exception as e:
            raise ValueError(f"Failed to open DOCX document: {e}") from e

        # Group findings by target container
        # 1. Body paragraphs: keyed by paragraph_index
        p_findings: Dict[int, List[Finding]] = {}
        # 2. Table cells: keyed by (tbl_idx, r_idx, c_idx)
        table_findings: Dict[Tuple[int, int, int], List[Finding]] = {}
        # 3. Unassigned / fallback
        unassigned: List[Finding] = []

        for f in to_redact:
            assigned = False
            if f.location and f.location.paragraph_index is not None:
                p_idx = f.location.paragraph_index
                if 0 <= p_idx < len(doc.paragraphs):
                    p_findings.setdefault(p_idx, []).append(f)
                    assigned = True
            elif f.location and f.location.sheet_name and f.location.row and f.location.col:
                try:
                    tbl_idx = int(f.location.sheet_name.replace("Table", "")) - 1
                    r_idx = f.location.row - 1
                    c_idx = f.location.col - 1
                    if 0 <= tbl_idx < len(doc.tables):
                        table = doc.tables[tbl_idx]
                        if 0 <= r_idx < len(table.rows) and 0 <= c_idx < len(table.rows[r_idx].cells):
                            table_findings.setdefault((tbl_idx, r_idx, c_idx), []).append(f)
                            assigned = True
                except Exception:
                    pass

            if not assigned:
                unassigned.append(f)

        # 1. Redact body paragraphs with descending start_char sorting
        for p_idx, f_list in p_findings.items():
            p = doc.paragraphs[p_idx]

            def get_p_sort_key(item: Finding) -> int:
                if item.location and item.location.start_char is not None:
                    return item.location.start_char
                pos = p.text.find(item.matched_text)
                return pos if pos != -1 else 0

            f_list.sort(key=get_p_sort_key, reverse=True)
            for f in f_list:
                s_char = f.location.start_char if f.location else None
                e_char = f.location.end_char if f.location else None
                _redact_in_paragraph(p, f.matched_text, s_char, e_char)

        # 2. Redact table cells with descending start_char sorting
        for (tbl_idx, r_idx, c_idx), f_list in table_findings.items():
            table = doc.tables[tbl_idx]
            cell = table.cell(r_idx, c_idx)

            def get_t_sort_key(item: Finding) -> int:
                if item.location and item.location.start_char is not None:
                    return item.location.start_char
                pos = cell.text.find(item.matched_text)
                return pos if pos != -1 else 0

            f_list.sort(key=get_t_sort_key, reverse=True)
            for f in f_list:
                s_char = f.location.start_char if f.location else None
                e_char = f.location.end_char if f.location else None
                redacted = False
                for p in cell.paragraphs:
                    if f.matched_text in p.text:
                        if _redact_in_paragraph(p, f.matched_text, s_char, e_char):
                            redacted = True
                            break
                if not redacted:
                    unassigned.append(f)

        # 3. Fallback for unassigned findings
        for f in unassigned:
            redacted = False
            s_char = f.location.start_char if f.location else None
            e_char = f.location.end_char if f.location else None
            for p in doc.paragraphs:
                if f.matched_text in p.text:
                    if _redact_in_paragraph(p, f.matched_text, s_char, e_char):
                        redacted = True
                        break
            if not redacted:
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                if f.matched_text in p.text:
                                    if _redact_in_paragraph(p, f.matched_text, s_char, e_char):
                                        redacted = True
                                        break
                                if redacted:
                                    break
                            if redacted:
                                break
                        if redacted:
                            break

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
