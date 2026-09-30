"""DOCX document parser using python-docx for paragraphs, tables, and character offsets."""
import io
from typing import List
import docx
from app.models.schemas import Finding, Location
from app.core.presidio_engine import get_presidio_engine


class DocxParser:
    """Extracts text and structural coordinates (paragraphs, tables) from Word documents."""

    def __init__(self):
        self.engine = get_presidio_engine()

    def parse_and_scan(self, content: bytes, file_name: str, start_id: int = 1) -> List[Finding]:
        """Scan DOCX document and return findings mapped to paragraph or table coordinates."""
        findings: List[Finding] = []
        curr_id = start_id

        try:
            doc = docx.Document(io.BytesIO(content))
        except Exception as e:
            raise ValueError(f"Failed to open DOCX file: {e}") from e

        # 1. Scan body paragraphs
        for p_idx, paragraph in enumerate(doc.paragraphs):
            p_text = paragraph.text
            if not p_text or not p_text.strip():
                continue

            detected_items = self.engine.analyze_text(p_text)
            for item in detected_items:
                location = Location(
                    paragraph_index=p_idx,
                    start_char=item.start,
                    end_char=item.end,
                )
                findings.append(
                    Finding(
                        id=f"find_{curr_id}",
                        entity_type=item.entity_type,
                        matched_text=item.matched_text,
                        confidence=item.confidence,
                        reasoning=item.reasoning,
                        location=location,
                    )
                )
                curr_id += 1

        # 2. Scan table cells
        for tbl_idx, table in enumerate(doc.tables):
            for row_idx, row in enumerate(table.rows, start=1):
                for col_idx, cell in enumerate(row.cells, start=1):
                    cell_text = cell.text
                    if not cell_text or not cell_text.strip():
                        continue

                    detected_items = self.engine.analyze_text(cell_text)
                    for item in detected_items:
                        location = Location(
                            sheet_name=f"Table{tbl_idx + 1}",
                            cell=f"T{tbl_idx + 1}_R{row_idx}C{col_idx}",
                            row=row_idx,
                            col=col_idx,
                            start_char=item.start,
                            end_char=item.end,
                        )
                        findings.append(
                            Finding(
                                id=f"find_{curr_id}",
                                entity_type=item.entity_type,
                                matched_text=item.matched_text,
                                confidence=item.confidence,
                                reasoning=item.reasoning,
                                location=location,
                            )
                        )
                        curr_id += 1

        return findings
