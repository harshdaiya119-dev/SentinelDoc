"""Spreadsheet parser for XLSX and CSV files with cell coordinate mapping."""
import io
import csv
from typing import List
import openpyxl
from openpyxl.utils import get_column_letter
from app.models.schemas import Finding, Location
from app.core.presidio_engine import get_presidio_engine


class SpreadsheetParser:
    """Extracts text and cell coordinates (sheet, row, col, cell reference) from XLSX and CSV."""

    def __init__(self):
        self.engine = get_presidio_engine()

    def parse_and_scan(self, content: bytes, file_name: str, start_id: int = 1) -> List[Finding]:
        """Dispatch parsing based on file extension (.xlsx vs .csv)."""
        lower_name = file_name.lower()
        if lower_name.endswith(".csv"):
            return self._parse_csv(content, start_id)
        elif lower_name.endswith(".xlsx"):
            return self._parse_xlsx(content, start_id)
        else:
            raise ValueError(f"Unsupported spreadsheet format: {file_name}")

    def _parse_xlsx(self, content: bytes, start_id: int) -> List[Finding]:
        """Scan Excel workbook cells using openpyxl."""
        findings: List[Finding] = []
        curr_id = start_id

        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        except Exception as e:
            raise ValueError(f"Failed to open XLSX workbook: {e}") from e

        try:
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                for row in sheet.iter_rows(values_only=False):
                    for cell in row:
                        if cell.value is None:
                            continue
                        cell_str = str(cell.value)
                        if not cell_str.strip():
                            continue

                        detected_items = self.engine.analyze_text(cell_str)
                        for item in detected_items:
                            location = Location(
                                sheet_name=sheet_name,
                                cell=cell.coordinate,
                                row=cell.row,
                                col=cell.column,
                                start_char=item.start,
                                end_char=item.end,
                            )
                            findings.append(
                                Finding(
                                    id=f"find_{curr_id}",
                                    entity_type=item.entity_type,
                                    matched_text=item.matched_text,
                                    confidence=item.confidence,
                                    reasoning=item.reasoning,
                                    location=location,
                                )
                            )
                            curr_id += 1
        finally:
            wb.close()

        return findings

    def _parse_csv(self, content: bytes, start_id: int) -> List[Finding]:
        """Scan CSV rows and columns using csv reader."""
        findings: List[Finding] = []
        curr_id = start_id

        # Try UTF-8 decoding with fallback
        try:
            text_data = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text_data = content.decode("latin-1")
            except Exception as e:
                raise ValueError(f"Failed to decode CSV content: {e}") from e

        reader = csv.reader(io.StringIO(text_data))
        for row_idx, row in enumerate(reader, start=1):
            for col_idx, cell_value in enumerate(row, start=1):
                if not cell_value or not cell_value.strip():
                    continue

                col_letter = get_column_letter(col_idx)
                cell_coord = f"{col_letter}{row_idx}"

                detected_items = self.engine.analyze_text(cell_value)
                for item in detected_items:
                    location = Location(
                        sheet_name="Sheet1",
                        cell=cell_coord,
                        row=row_idx,
                        col=col_idx,
                        start_char=item.start,
                        end_char=item.end,
                    )
                    findings.append(
                        Finding(
                            id=f"find_{curr_id}",
                            entity_type=item.entity_type,
                            matched_text=item.matched_text,
                            confidence=item.confidence,
                            reasoning=item.reasoning,
                            location=location,
                        )
                    )
                    curr_id += 1

        return findings
