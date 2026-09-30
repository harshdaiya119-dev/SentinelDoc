"""Adversarial stress and coordinate boundary test suite for POST /scan.

Authored by teamwork_preview_challenger_m1_2 (Stress & Coordinate Challenger).
Empirically stress-tests:
1. Multi-page PDFs and bounding box accuracy (x0, top, x1, bottom, duplicate occurrence distinction).
2. DOCX paragraphs, character offsets, and multi-table cell coordinates.
3. XLSX multi-sheet isolation, numeric cells, and column coordinates beyond Z.
4. CSV quoted commas, multiline cells, and UTF-8 BOM encoding.
5. TXT multi-line mapping, CRLF line endings, and Unicode offset alignment.
6. Edge conditions: 0-byte files, unsupported extensions, corrupted magic bytes, truncated zip/pdf streams.
"""
import io
import csv
import docx
import fitz
import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app
from tests.fixtures.generator import (
    create_txt,
    create_csv,
    create_xlsx,
    create_docx,
    create_pdf,
    generate_simple_pdf,
    generate_simple_docx,
    generate_simple_xlsx,
    generate_simple_csv,
)


@pytest.fixture(scope="module")
def client() -> TestClient:
    return TestClient(app)


# =====================================================================
# 1. PDF STRESS & COORDINATE BOUNDARY TESTS
# =====================================================================

class TestPdfStressAndCoordinates:
    """Adversarial coordinate and multi-page stress tests for PDF documents."""

    def test_pdf_multipage_page_isolation(self, client: TestClient):
        """Verify findings across multi-page PDF map to exact 1-indexed pages and empty pages are skipped."""
        doc = fitz.open()
        
        # Page 1: Aadhaar
        p1 = doc.new_page(width=612, height=792)
        p1.insert_text(fitz.Point(72, 100), "Identity Record\nAadhaar: 3675 9832 4511")

        # Page 2: Blank page (no text)
        doc.new_page(width=612, height=792)

        # Page 3: PAN and Card
        p3 = doc.new_page(width=612, height=792)
        p3.insert_text(fitz.Point(72, 120), "Financial Identifiers\nPAN: ABCPE1234F\nCard: 4532 0151 1283 0366")

        # Page 4: Unrelated filler without PII
        p4 = doc.new_page(width=612, height=792)
        p4.insert_text(fitz.Point(72, 100), "Corporate policy and governance memo without confidential identifiers.")

        # Page 5: Contact phone and email
        p5 = doc.new_page(width=612, height=792)
        p5.insert_text(fitz.Point(72, 150), "Support contact: +91 98765 43210\nOfficial email: security@sentineldoc.com")

        pdf_bytes = doc.tobytes()
        doc.close()

        resp = client.post("/scan", files={"file": ("multipage_test.pdf", pdf_bytes, "application/pdf")})
        assert resp.status_code == 200, f"Expected 200 OK, got {resp.status_code}: {resp.text}"
        data = resp.json()
        assert data["file_type"] == "pdf"
        findings = data["findings"]

        # Collect findings by page
        page_map = {}
        for f in findings:
            loc = f["location"]
            page = loc["page"]
            page_map.setdefault(page, []).append(f)

        # Verify page mapping
        assert 1 in page_map, "Expected findings on Page 1"
        assert 3 in page_map, "Expected findings on Page 3"
        assert 5 in page_map, "Expected findings on Page 5"
        assert 2 not in page_map, "No findings should exist on blank Page 2"
        assert 4 not in page_map, "No findings should exist on non-PII Page 4"

        # Check Page 1 has Aadhaar
        p1_types = [f["entity_type"] for f in page_map[1]]
        assert "AADHAAR" in p1_types

        # Check Page 3 has PAN and Card
        p3_types = [f["entity_type"] for f in page_map[3]]
        assert "PAN" in p3_types
        assert "CREDIT_CARD" in p3_types

        # Check Page 5 has Phone and Email
        p5_types = [f["entity_type"] for f in page_map[5]]
        assert "PHONE_NUMBER" in p5_types
        assert "EMAIL_ADDRESS" in p5_types

    def test_pdf_bounding_box_accuracy(self, client: TestClient):
        """Verify bounding box coordinates (x0, top, x1, bottom) are strictly valid and non-degenerate."""
        doc = fitz.open()
        page = doc.new_page(width=612, height=792)
        # Place text at known position
        page.insert_text(fitz.Point(100.0, 200.0), "Planted Aadhaar: 3675 9832 4511", fontname="helv", fontsize=12)
        page.insert_text(fitz.Point(100.0, 400.0), "Planted PAN: ABCPE1234F", fontname="helv", fontsize=12)
        pdf_bytes = doc.tobytes()
        doc.close()

        resp = client.post("/scan", files={"file": ("coords_test.pdf", pdf_bytes, "application/pdf")})
        assert resp.status_code == 200
        findings = resp.json()["findings"]
        assert len(findings) >= 2

        for f in findings:
            loc = f["location"]
            assert loc["page"] == 1
            bbox = loc["bbox"]
            assert bbox is not None, f"Finding {f['id']} missing bbox"

            x0, top, x1, bottom = bbox["x0"], bbox["top"], bbox["x1"], bbox["bottom"]
            # Verify coordinates are positive floats
            assert isinstance(x0, float) and isinstance(top, float)
            assert isinstance(x1, float) and isinstance(bottom, float)
            
            # Geometry checks: positive width and height
            assert x0 < x1, f"Invalid horizontal bbox geometry: x0={x0} >= x1={x1}"
            assert top < bottom, f"Invalid vertical bbox geometry: top={top} >= bottom={bottom}"

            # Ensure coordinates are within page dimensions (612 x 792)
            assert 0.0 <= x0 <= 612.0
            assert 0.0 <= x1 <= 612.0
            assert 0.0 <= top <= 792.0
            assert 0.0 <= bottom <= 792.0

        # Specific coordinate placement checks
        aadhaar_f = next(f for f in findings if f["entity_type"] == "AADHAAR")
        abox = aadhaar_f["location"]["bbox"]
        # y=200 text baseline -> top should be around ~188-202, bottom around ~200-205
        assert 180.0 <= abox["top"] <= 210.0, f"Aadhaar top {abox['top']} outside expected vertical zone"

        pan_f = next(f for f in findings if f["entity_type"] == "PAN")
        pbox = pan_f["location"]["bbox"]
        assert 380.0 <= pbox["top"] <= 410.0, f"PAN top {pbox['top']} outside expected vertical zone"

    def test_pdf_duplicate_pii_different_bounding_boxes(self, client: TestClient):
        """Verify identical PII string repeated 3 times on the same page produces 3 distinct vertical bboxes."""
        doc = fitz.open()
        page = doc.new_page(width=612, height=792)
        page.insert_text(fitz.Point(72.0, 100.0), "Section 1 Aadhaar: 3675 9832 4511", fontname="helv", fontsize=11)
        page.insert_text(fitz.Point(72.0, 350.0), "Section 2 Aadhaar: 3675 9832 4511", fontname="helv", fontsize=11)
        page.insert_text(fitz.Point(72.0, 600.0), "Section 3 Aadhaar: 3675 9832 4511", fontname="helv", fontsize=11)
        pdf_bytes = doc.tobytes()
        doc.close()

        resp = client.post("/scan", files={"file": ("duplicate_coords.pdf", pdf_bytes, "application/pdf")})
        assert resp.status_code == 200
        findings = [f for f in resp.json()["findings"] if f["entity_type"] == "AADHAAR"]
        assert len(findings) == 3, f"Expected 3 Aadhaar findings, got {len(findings)}"

        tops = [f["location"]["bbox"]["top"] for f in findings]
        # Verify that each top coordinate is distinct and corresponds to the distinct text lines
        assert len(set(tops)) == 3, f"Expected 3 distinct vertical coordinates for occurrences, got tops={tops}"
        assert abs(tops[0] - 100.0) < 25.0
        assert abs(tops[1] - 350.0) < 25.0
        assert abs(tops[2] - 600.0) < 25.0


# =====================================================================
# 2. DOCX PARAGRAPH & TABLE TESTS
# =====================================================================

class TestDocxStressAndCoordinates:
    """Stress tests for Word (.docx) documents verifying paragraphs, runs, and tables."""

    def test_docx_paragraph_indices_and_char_offsets(self, client: TestClient):
        """Verify paragraph indices and start/end character offsets match original text."""
        doc = docx.Document()
        doc.add_paragraph("Document Header - SentinelDoc Compliance")
        doc.add_paragraph("    ")  # Empty paragraph with spaces
        doc.add_paragraph("Employee record for internal audit purposes.")
        p3 = doc.add_paragraph("Tax identifier PAN: ABCPE1234F registered with NSDL.")
        doc.add_paragraph("")  # Completely empty paragraph
        p5 = doc.add_paragraph("Primary Aadhaar reference 3675 9832 4511.")
        
        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        resp = client.post("/scan", files={"file": ("paragraphs.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert resp.status_code == 200
        findings = resp.json()["findings"]

        pan_findings = [f for f in findings if f["entity_type"] == "PAN"]
        assert len(pan_findings) == 1
        pan_f = pan_findings[0]
        assert pan_f["location"]["paragraph_index"] == 3
        # Verify character offset extraction slices PAN exactly
        p3_text = p3.text
        start, end = pan_f["location"]["start_char"], pan_f["location"]["end_char"]
        assert p3_text[start:end] == "ABCPE1234F"

        aadhaar_findings = [f for f in findings if f["entity_type"] == "AADHAAR"]
        assert len(aadhaar_findings) == 1
        aadhaar_f = aadhaar_findings[0]
        assert aadhaar_f["location"]["paragraph_index"] == 5
        p5_text = p5.text
        start, end = aadhaar_f["location"]["start_char"], aadhaar_f["location"]["end_char"]
        assert p5_text[start:end] == "3675 9832 4511"

    def test_docx_multiple_tables_and_cell_coordinates(self, client: TestClient):
        """Verify multiple tables extract correct sheet_name, cell reference, row, col, and char offsets."""
        doc = docx.Document()
        doc.add_paragraph("Executive Summary with Embedded Data Tables")

        # Table 1: 3 rows, 3 columns
        t1 = doc.add_table(rows=3, cols=3)
        t1.cell(0, 0).text = "User ID"
        t1.cell(0, 1).text = "Full Name"
        t1.cell(0, 2).text = "Identity Document"
        t1.cell(1, 0).text = "USR-001"
        t1.cell(1, 1).text = "Vikram Sharma"
        t1.cell(1, 2).text = "Aadhaar: 3675 9832 4511"
        t1.cell(2, 0).text = "USR-002"
        t1.cell(2, 1).text = "Priya Rao"
        t1.cell(2, 2).text = "PAN: XYZCA5678B"

        doc.add_paragraph("Payment and billing table follows:")

        # Table 2: 2 rows, 2 columns
        t2 = doc.add_table(rows=2, cols=2)
        t2.cell(0, 0).text = "Account Type"
        t2.cell(0, 1).text = "Card Number"
        t2.cell(1, 0).text = "Corporate Visa"
        t2.cell(1, 1).text = "4532 0151 1283 0366"

        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        resp = client.post("/scan", files={"file": ("tables.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert resp.status_code == 200
        findings = resp.json()["findings"]

        # Check Table 1 Aadhaar
        aadhaar_f = next(f for f in findings if f["entity_type"] == "AADHAAR")
        loc1 = aadhaar_f["location"]
        assert loc1["sheet_name"] == "Table1"
        assert loc1["cell"] == "T1_R2C3"
        assert loc1["row"] == 2
        assert loc1["col"] == 3
        # Verify offset inside cell text
        cell1_text = t1.cell(1, 2).text
        assert cell1_text[loc1["start_char"]:loc1["end_char"]] == "3675 9832 4511"

        # Check Table 1 PAN
        pan_f = next(f for f in findings if f["entity_type"] == "PAN")
        loc2 = pan_f["location"]
        assert loc2["sheet_name"] == "Table1"
        assert loc2["cell"] == "T1_R3C3"
        assert loc2["row"] == 3
        assert loc2["col"] == 3
        cell2_text = t1.cell(2, 2).text
        assert cell2_text[loc2["start_char"]:loc2["end_char"]] == "XYZCA5678B"

        # Check Table 2 Card
        card_f = next(f for f in findings if f["entity_type"] == "CREDIT_CARD")
        loc3 = card_f["location"]
        assert loc3["sheet_name"] == "Table2"
        assert loc3["cell"] == "T2_R2C2"
        assert loc3["row"] == 2
        assert loc3["col"] == 2
        cell3_text = t2.cell(1, 1).text
        assert cell3_text[loc3["start_char"]:loc3["end_char"]] == "4532 0151 1283 0366"


# =====================================================================
# 3. XLSX MULTI-SHEET & NUMERIC CELLS TESTS
# =====================================================================

class TestXlsxStressAndCoordinates:
    """Stress tests for Excel (.xlsx) workbooks verifying multi-sheet isolation and numeric cell types."""

    def test_xlsx_multisheet_isolation(self, client: TestClient):
        """Verify findings correctly record their respective sheet names in a multi-sheet workbook."""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default sheet

        # Sheet 1: HR
        ws_hr = wb.create_sheet(title="HR_Records")
        ws_hr["A1"] = "Employee"
        ws_hr["B1"] = "Aadhaar Number"
        ws_hr["A2"] = "Aarav Sharma"
        ws_hr["B2"] = "3675 9832 4511"

        # Sheet 2: Empty sheet
        wb.create_sheet(title="Blank_Notes")

        # Sheet 3: Finance
        ws_fin = wb.create_sheet(title="Finance_Ledger")
        ws_fin["A1"] = "Account"
        ws_fin["B1"] = "Tax PAN"
        ws_fin["C1"] = "Payment Card"
        ws_fin["A2"] = "Treasury"
        ws_fin["B2"] = "ABCPE1234F"
        ws_fin["C2"] = "4532 0151 1283 0366"

        buf = io.BytesIO()
        wb.save(buf)
        xlsx_bytes = buf.getvalue()

        resp = client.post("/scan", files={"file": ("multisheet.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert resp.status_code == 200
        findings = resp.json()["findings"]

        # Map findings by sheet name
        sheet_findings = {}
        for f in findings:
            s_name = f["location"]["sheet_name"]
            sheet_findings.setdefault(s_name, []).append(f)

        assert "HR_Records" in sheet_findings
        assert "Finance_Ledger" in sheet_findings
        assert "Blank_Notes" not in sheet_findings

        hr_aadhaar = next(f for f in sheet_findings["HR_Records"] if f["entity_type"] == "AADHAAR")
        assert hr_aadhaar["location"]["cell"] == "B2"
        assert hr_aadhaar["location"]["row"] == 2
        assert hr_aadhaar["location"]["col"] == 2

        fin_pan = next(f for f in sheet_findings["Finance_Ledger"] if f["entity_type"] == "PAN")
        assert fin_pan["location"]["cell"] == "B2"

        fin_card = next(f for f in sheet_findings["Finance_Ledger"] if f["entity_type"] == "CREDIT_CARD")
        assert fin_card["location"]["cell"] == "C2"

    def test_xlsx_numeric_cells_and_extreme_columns(self, client: TestClient):
        """Verify numeric cell types (stored as int in Excel) and columns beyond Z (e.g. AA, AB)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "NumericData"

        # Stored as raw integer (not string!)
        ws["A1"] = 367598324511  # 12-digit integer Aadhaar
        ws["B1"] = 4532015112830366  # 16-digit integer Card

        # Harmless financial floats and order IDs (must NOT be flagged)
        ws["C1"] = 1234567.89  # Currency float
        ws["D1"] = 123456789012  # 12-digit number starting with 1 (fails Verhoeff / not Aadhaar)

        # Place a finding far out in column AA (Col 27) and AB (Col 28)
        ws.cell(row=5, column=27, value="ABCPE1234F")
        ws.cell(row=5, column=28, value="+91 98765 43210")

        buf = io.BytesIO()
        wb.save(buf)
        xlsx_bytes = buf.getvalue()

        resp = client.post("/scan", files={"file": ("numeric_extreme.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert resp.status_code == 200
        findings = resp.json()["findings"]

        # Verify numeric unspaced Aadhaar was detected
        aadhaar_finds = [f for f in findings if f["entity_type"] == "AADHAAR"]
        assert len(aadhaar_finds) >= 1
        assert aadhaar_finds[0]["location"]["cell"] == "A1"

        # Verify numeric unspaced Card was detected
        card_finds = [f for f in findings if f["entity_type"] == "CREDIT_CARD"]
        assert len(card_finds) >= 1
        assert card_finds[0]["location"]["cell"] == "B1"

        # Verify false positives on C1 and D1 were rejected
        assert not any(f["location"]["cell"] == "C1" for f in findings)
        assert not any(f["location"]["cell"] == "D1" for f in findings)

        # Verify column AA5 and AB5 coordinates
        pan_aa = next(f for f in findings if f["entity_type"] == "PAN")
        assert pan_aa["location"]["cell"] == "AA5"
        assert pan_aa["location"]["row"] == 5
        assert pan_aa["location"]["col"] == 27

        phone_ab = next(f for f in findings if f["entity_type"] == "PHONE_NUMBER")
        assert phone_ab["location"]["cell"] == "AB5"
        assert phone_ab["location"]["row"] == 5
        assert phone_ab["location"]["col"] == 28


# =====================================================================
# 4. CSV QUOTED COMMAS & MULTILINE TESTS
# =====================================================================

class TestCsvStressAndCoordinates:
    """Stress tests for CSV documents verifying RFC 4180 quoted commas, multiline cells, and encoding."""

    def test_csv_quoted_commas_column_integrity(self, client: TestClient):
        """Verify cells containing quoted commas do not cause column shift in subsequent fields."""
        # Col 1: ID, Col 2: Name with comma, Col 3: Address with commas, Col 4: PAN
        csv_text = (
            'id,name,address,pan\n'
            '1,"Sharma, Rajesh","Flat 101, Complex B, MG Road, Bangalore",ABCPE1234F\n'
            '2,"Mukherjee, Vikram","24B, Park Street, Kolkata",XYZCA5678B\n'
        )
        resp = client.post("/scan", files={"file": ("quoted.csv", csv_text.encode("utf-8"), "text/csv")})
        assert resp.status_code == 200
        findings = resp.json()["findings"]

        pan_findings = [f for f in findings if f["entity_type"] == "PAN"]
        assert len(pan_findings) == 2

        # Row 2 (index 2): PAN must be in Column 4 (D2)
        pan_r2 = next(f for f in pan_findings if f["location"]["row"] == 2)
        assert pan_r2["location"]["col"] == 4, f"Column shifted! Expected col 4 (D), got {pan_r2['location']['col']}"
        assert pan_r2["location"]["cell"] == "D2"

        # Row 3 (index 3): PAN must be in Column 4 (D3)
        pan_r3 = next(f for f in pan_findings if f["location"]["row"] == 3)
        assert pan_r3["location"]["col"] == 4
        assert pan_r3["location"]["cell"] == "D3"

    def test_csv_multiline_cells(self, client: TestClient):
        """Verify embedded newlines within quoted CSV cells do not break row numbering."""
        csv_text = (
            'id,notes,contact\n'
            '1,"Customer statement:\nLine 2 info\nAadhaar: 3675 9832 4511\nEnd note",+91 98765 43210\n'
            '2,"Simple note",user@sentineldoc.com\n'
        )
        resp = client.post("/scan", files={"file": ("multiline.csv", csv_text.encode("utf-8"), "text/csv")})
        assert resp.status_code == 200
        findings = resp.json()["findings"]

        # Aadhaar is in row 2, col 2 (B2)
        aadhaar_f = next(f for f in findings if f["entity_type"] == "AADHAAR")
        assert aadhaar_f["location"]["row"] == 2
        assert aadhaar_f["location"]["col"] == 2
        assert aadhaar_f["location"]["cell"] == "B2"

        # Phone is in row 2, col 3 (C2)
        phone_f = next(f for f in findings if f["entity_type"] == "PHONE_NUMBER")
        assert phone_f["location"]["row"] == 2
        assert phone_f["location"]["col"] == 3
        assert phone_f["location"]["cell"] == "C2"

        # Email is in row 3, col 3 (C3)
        email_f = next(f for f in findings if f["entity_type"] == "EMAIL_ADDRESS")
        assert email_f["location"]["row"] == 3
        assert email_f["location"]["col"] == 3
        assert email_f["location"]["cell"] == "C3"

    def test_csv_utf8_bom_encoding(self, client: TestClient):
        """Verify CSV files with UTF-8 BOM (\xef\xbb\xbf) are decoded cleanly without corrupting cell A1."""
        csv_bytes = b'\xef\xbb\xbfPAN,Card\nABCPE1234F,4532 0151 1283 0366\n'
        resp = client.post("/scan", files={"file": ("bom.csv", csv_bytes, "text/csv")})
        assert resp.status_code == 200
        findings = resp.json()["findings"]

        pan_f = next(f for f in findings if f["entity_type"] == "PAN")
        assert pan_f["location"]["cell"] == "A2"
        assert pan_f["matched_text"] == "ABCPE1234F"


# =====================================================================
# 5. TXT MULTI-LINE & CHARACTER OFFSET TESTS
# =====================================================================

class TestTxtStressAndCoordinates:
    """Stress tests for plain text files verifying line numbering, CRLF endings, and char offsets."""

    def test_txt_multiline_character_offsets(self, client: TestClient):
        """Verify exact character offsets and line numbers across a 100-line text document."""
        lines = [f"System log event {i}: Normal operational status" for i in range(1, 101)]
        # Plant PII on specific lines
        lines[19] = "Audit Alert: Planted Aadhaar is 3675 9832 4511 verified."  # Line 20
        lines[49] = "Compliance Check: Taxpayer PAN ABCPE1234F verified."         # Line 50
        lines[79] = "Billing Notice: Card 4532 0151 1283 0366 debited."          # Line 80
        lines[99] = "Helpdesk Phone: +91 98765 43210 available."                  # Line 100

        txt_content = "\n".join(lines)
        resp = client.post("/scan", files={"file": ("audit_100.txt", txt_content.encode("utf-8"), "text/plain")})
        assert resp.status_code == 200
        findings = resp.json()["findings"]

        # Aadhaar on Line 20
        aadhaar_f = next(f for f in findings if f["entity_type"] == "AADHAAR")
        assert aadhaar_f["location"]["line"] == 20
        start, end = aadhaar_f["location"]["start_char"], aadhaar_f["location"]["end_char"]
        assert lines[19][start:end] == "3675 9832 4511"

        # PAN on Line 50
        pan_f = next(f for f in findings if f["entity_type"] == "PAN")
        assert pan_f["location"]["line"] == 50
        start, end = pan_f["location"]["start_char"], pan_f["location"]["end_char"]
        assert lines[49][start:end] == "ABCPE1234F"

        # Card on Line 80
        card_f = next(f for f in findings if f["entity_type"] == "CREDIT_CARD")
        assert card_f["location"]["line"] == 80
        start, end = card_f["location"]["start_char"], card_f["location"]["end_char"]
        assert lines[79][start:end] == "4532 0151 1283 0366"

        # Phone on Line 100
        phone_f = next(f for f in findings if f["entity_type"] == "PHONE_NUMBER")
        assert phone_f["location"]["line"] == 100
        start, end = phone_f["location"]["start_char"], phone_f["location"]["end_char"]
        assert lines[99][start:end] == "+91 98765 43210"

    def test_txt_crlf_windows_newlines(self, client: TestClient):
        """Verify Windows CRLF (\\r\\n) line endings preserve exact line indices and offsets."""
        lines = [
            "Header line with no sensitive data",
            "Second line with PAN: ABCPE1234F",
            "Third line with Card: 4532 0151 1283 0366",
            "Fourth line ending",
        ]
        crlf_content = "\r\n".join(lines) + "\r\n"
        resp = client.post("/scan", files={"file": ("windows_crlf.txt", crlf_content.encode("utf-8"), "text/plain")})
        assert resp.status_code == 200
        findings = resp.json()["findings"]

        pan_f = next(f for f in findings if f["entity_type"] == "PAN")
        assert pan_f["location"]["line"] == 2
        start, end = pan_f["location"]["start_char"], pan_f["location"]["end_char"]
        assert lines[1][start:end] == "ABCPE1234F"


# =====================================================================
# 6. EDGE CONDITIONS: 0-BYTE, UNSUPPORTED, CORRUPTED MAGIC BYTES
# =====================================================================

class TestEdgeConditionsAndMalformedUploads:
    """Rigorous boundary and vulnerability tests on degenerate inputs."""

    @pytest.mark.parametrize("ext", [".pdf", ".docx", ".xlsx", ".csv", ".txt"])
    def test_0_byte_empty_files_return_400_empty_file(self, client: TestClient, ext: str):
        """Verify 0-byte upload across all 5 formats returns HTTP 400 with code EMPTY_FILE."""
        filename = f"empty_test{ext}"
        resp = client.post("/scan", files={"file": (filename, b"", "application/octet-stream")})
        assert resp.status_code == 400, f"Expected 400 for 0-byte {ext}, got {resp.status_code}: {resp.text}"
        data = resp.json()
        assert "error" in data
        assert data["error"]["code"] == "EMPTY_FILE"
        assert data["error"]["status_code"] == 400

    @pytest.mark.parametrize("ext", [".exe", ".bin", ".py", ".sh", ".zip", ".png", ".jpg", ".odt", ".html"])
    def test_unsupported_extensions_return_400_unsupported_file_type(self, client: TestClient, ext: str):
        """Verify unsupported file extensions return HTTP 400 with code UNSUPPORTED_FILE_TYPE."""
        filename = f"unsupported_payload{ext}"
        resp = client.post("/scan", files={"file": (filename, b"dummy payload content", "application/octet-stream")})
        assert resp.status_code == 400
        data = resp.json()
        assert "error" in data
        assert data["error"]["code"] == "UNSUPPORTED_FILE_TYPE"
        assert data["error"]["status_code"] == 400

    def test_corrupted_pdf_missing_magic_header(self, client: TestClient):
        """Verify PDF missing %PDF- header returns HTTP 400 with code MALFORMED_FILE."""
        garbage = b"This is not a PDF file at all, just random plain text."
        resp = client.post("/scan", files={"file": ("fake.pdf", garbage, "application/pdf")})
        assert resp.status_code == 400
        data = resp.json()
        assert data["error"]["code"] == "MALFORMED_FILE"
        assert data["error"]["status_code"] == 400

    def test_corrupted_docx_missing_zip_header(self, client: TestClient):
        """Verify DOCX missing PK\\x03\\x04 header returns HTTP 400 with code MALFORMED_FILE."""
        garbage = b"This is not a DOCX zip archive at all."
        resp = client.post("/scan", files={"file": ("fake.docx", garbage, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert resp.status_code == 400
        data = resp.json()
        assert data["error"]["code"] == "MALFORMED_FILE"
        assert data["error"]["status_code"] == 400

    def test_corrupted_xlsx_missing_zip_header(self, client: TestClient):
        """Verify XLSX missing PK\\x03\\x04 header returns HTTP 400 with code MALFORMED_FILE."""
        garbage = b"This is not an XLSX zip archive at all."
        resp = client.post("/scan", files={"file": ("fake.xlsx", garbage, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert resp.status_code == 400
        data = resp.json()
        assert data["error"]["code"] == "MALFORMED_FILE"
        assert data["error"]["status_code"] == 400

    def test_truncated_corrupt_pdf_with_magic_bytes(self, client: TestClient):
        """Verify PDF with valid magic header %PDF-1.4 but unparsable corrupt data returns 400 (not 500)."""
        truncated_pdf = b"%PDF-1.4\nTRUNCATED_NON_PARSABLE_BYTES\xff\x00\xff"
        resp = client.post("/scan", files={"file": ("truncated.pdf", truncated_pdf, "application/pdf")})
        # Must return clean 400 error envelope, never unhandled 500
        assert resp.status_code == 400
        data = resp.json()
        assert data["error"]["code"] == "MALFORMED_FILE"
        assert data["error"]["status_code"] == 400

    def test_truncated_corrupt_docx_with_magic_bytes(self, client: TestClient):
        """Verify DOCX with valid magic header PK\\x03\\x04 but truncated corrupt zip returns 400 (not 500)."""
        truncated_docx = b"PK\x03\x04\x14\x00\x00\x00\x08\x00_CORRUPT_NON_ZIP_STREAM"
        resp = client.post("/scan", files={"file": ("truncated.docx", truncated_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert resp.status_code == 400
        data = resp.json()
        assert data["error"]["code"] == "MALFORMED_FILE"

    def test_truncated_corrupt_xlsx_with_magic_bytes(self, client: TestClient):
        """Verify XLSX with valid magic header PK\\x03\\x04 but truncated corrupt zip returns 400 (not 500)."""
        truncated_xlsx = b"PK\x03\x04\x14\x00\x00\x00\x08\x00_CORRUPT_NON_ZIP_STREAM"
        resp = client.post("/scan", files={"file": ("truncated.xlsx", truncated_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert resp.status_code == 400
        data = resp.json()
        assert data["error"]["code"] == "MALFORMED_FILE"
        assert data["error"]["status_code"] == 400

    @pytest.mark.parametrize("filename,ext", [
        ("report.PDF", "pdf"),
        ("memo.DocX", "docx"),
        ("sheet.XLSX", "xlsx"),
        ("data.CsV", "csv"),
        ("log.TxT", "txt"),
        ("company.quarterly.audit.pdf", "pdf"),
        ("spaces in file name test.docx", "docx"),
    ])
    def test_mixed_case_extensions_and_complex_filenames(self, client: TestClient, filename: str, ext: str):
        """Verify mixed-case extensions (.PDF, .DocX) and multi-dot filenames are handled gracefully."""
        if ext == "pdf":
            content = generate_simple_pdf("PAN: ABCPE1234F")
        elif ext == "docx":
            content = generate_simple_docx("PAN: ABCPE1234F")
        elif ext == "xlsx":
            content = generate_simple_xlsx("Sheet1", ["PAN"], [["ABCPE1234F"]])
        elif ext == "csv":
            content = generate_simple_csv(["PAN"], [["ABCPE1234F"]])
        else:
            content = create_txt("PAN: ABCPE1234F\n")

        resp = client.post("/scan", files={"file": (filename, content, "application/octet-stream")})
        assert resp.status_code == 200, f"Failed for {filename}: {resp.text}"
        data = resp.json()
        assert data["file_type"] == ext.lower()
        assert len(data["findings"]) >= 1

    def test_filename_missing_extension(self, client: TestClient):
        """Verify file without any extension returns HTTP 400 UNSUPPORTED_FILE_TYPE."""
        resp = client.post("/scan", files={"file": ("no_extension_file", b"some text content", "application/octet-stream")})
        assert resp.status_code == 400
        data = resp.json()
        assert data["error"]["code"] == "UNSUPPORTED_FILE_TYPE"

    def test_docx_multi_run_paragraph_pii_integrity(self, client: TestClient):
        """Verify PII split across multiple text runs in a DOCX paragraph is accurately detected."""
        doc = docx.Document()
        p = doc.add_paragraph()
        r1 = p.add_run("Customer PAN identifier: ")
        r2 = p.add_run("ABCPE")
        r2.bold = True
        r3 = p.add_run("1234F")
        r3.italic = True
        r4 = p.add_run(" verified by tax authority.")

        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        resp = client.post("/scan", files={"file": ("multirun.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert resp.status_code == 200
        findings = resp.json()["findings"]
        pan_f = next((f for f in findings if f["entity_type"] == "PAN"), None)
        assert pan_f is not None, "Failed to detect PAN split across styled runs!"
        assert pan_f["matched_text"] == "ABCPE1234F"

    def test_csv_escaped_quotes(self, client: TestClient):
        """Verify CSV fields with escaped double quotes ("") parse correctly without breaking column boundaries."""
        csv_text = (
            'id,comment,pan\n'
            '1,"He remarked ""Tax ID ABCPE1234F is valid"" in review",ABCPE1234F\n'
        )
        resp = client.post("/scan", files={"file": ("escaped_quotes.csv", csv_text.encode("utf-8"), "text/csv")})
        assert resp.status_code == 200
        findings = resp.json()["findings"]
        # Column 3 PAN must be detected at col 3 (C2)
        pan_col3 = [f for f in findings if f["entity_type"] == "PAN" and f["location"]["col"] == 3]
        assert len(pan_col3) >= 1
        assert pan_col3[0]["location"]["cell"] == "C2"

    def test_txt_boundary_offsets_start_and_end_of_line(self, client: TestClient):
        """Verify character offsets when PII starts at char 0 and ends at exact line length."""
        # Line 1: PII at start (char 0)
        # Line 2: PII at end
        line1 = "3675 9832 4511 is the Aadhaar number."
        line2 = "The registered PAN is ABCPE1234F"
        txt_content = f"{line1}\n{line2}\n"

        resp = client.post("/scan", files={"file": ("boundary_offsets.txt", txt_content.encode("utf-8"), "text/plain")})
        assert resp.status_code == 200
        findings = resp.json()["findings"]

        aadhaar_f = next(f for f in findings if f["entity_type"] == "AADHAAR")
        assert aadhaar_f["location"]["line"] == 1
        assert aadhaar_f["location"]["start_char"] == 0
        assert aadhaar_f["location"]["end_char"] == 14
        assert line1[0:14] == "3675 9832 4511"

        pan_f = next(f for f in findings if f["entity_type"] == "PAN")
        assert pan_f["location"]["line"] == 2
        assert pan_f["location"]["end_char"] == len(line2)
        assert line2[pan_f["location"]["start_char"]:pan_f["location"]["end_char"]] == "ABCPE1234F"

    def test_pdf_50_page_stress(self, client: TestClient):
        """Stress test: 50-page PDF with findings on page 25 and 50."""
        doc = fitz.open()
        for i in range(50):
            page = doc.new_page(width=612, height=792)
            if i == 24:
                page.insert_text(fitz.Point(100, 200), "Aadhaar: 3675 9832 4511")
            elif i == 49:
                page.insert_text(fitz.Point(100, 300), "PAN: ABCPE1234F")
            else:
                page.insert_text(fitz.Point(100, 100), f"Page {i+1} filler")
        pdf_bytes = doc.tobytes()
        doc.close()

        resp = client.post("/scan", files={"file": ("50page.pdf", pdf_bytes, "application/pdf")})
        assert resp.status_code == 200
        finds = resp.json()["findings"]
        assert len(finds) == 2
        assert finds[0]["location"]["page"] == 25
        assert finds[1]["location"]["page"] == 50

    def test_csv_jagged_rows_and_uneven_columns(self, client: TestClient):
        """Stress test: CSV with uneven row lengths (jagged columns) correctly resolves coordinates."""
        jagged_csv = b"col1,col2\n1,2,3,4\n5\n6,7,\"PAN: ABCPE1234F\",9,10,11\n"
        resp = client.post("/scan", files={"file": ("jagged.csv", jagged_csv, "text/csv")})
        assert resp.status_code == 200
        finds = resp.json()["findings"]
        assert len(finds) == 1
        assert finds[0]["location"]["row"] == 4
        assert finds[0]["location"]["col"] == 3
        assert finds[0]["location"]["cell"] == "C4"

    def test_docx_50_row_table_stress(self, client: TestClient):
        """Stress test: DOCX table with 50 rows, with PII placed at row 45."""
        doc = docx.Document()
        tbl = doc.add_table(rows=50, cols=3)
        for r in range(50):
            tbl.cell(r, 0).text = f"ID-{r}"
            tbl.cell(r, 1).text = f"User-{r}"
            if r == 44:
                tbl.cell(r, 2).text = "Aadhaar: 3675 9832 4511"
            else:
                tbl.cell(r, 2).text = "Standard entry"
        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        resp = client.post("/scan", files={"file": ("table_50rows.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert resp.status_code == 200
        finds = resp.json()["findings"]
        assert len(finds) >= 1
        aadhaar_f = next(f for f in finds if f["entity_type"] == "AADHAAR")
        assert aadhaar_f["location"]["row"] == 45
        assert aadhaar_f["location"]["col"] == 3
        assert aadhaar_f["location"]["cell"] == "T1_R45C3"

    def test_txt_5000_lines_stress(self, client: TestClient):
        """Stress test: TXT document with 5,000 lines, PII planted at line 4,872."""
        lines = [f"Operational log line {i}" for i in range(1, 5001)]
        lines[4871] = "Security alert: Card 4532 0151 1283 0366 charged."  # Line 4872
        txt_content = "\n".join(lines)

        resp = client.post("/scan", files={"file": ("large_5000.txt", txt_content.encode("utf-8"), "text/plain")})
        assert resp.status_code == 200
        finds = resp.json()["findings"]
        card_f = next(f for f in finds if f["entity_type"] == "CREDIT_CARD")
        assert card_f["location"]["line"] == 4872
        start, end = card_f["location"]["start_char"], card_f["location"]["end_char"]
        assert lines[4871][start:end] == "4532 0151 1283 0366"


