"""Programmatic document generator synthesizing valid PDF, DOCX, CSV, XLSX, and TXT files.

Provides deterministic test document synthesis with planted PII and ground-truth metadata.
"""
import io
import csv
from typing import List, Dict, Tuple, Any, Optional
import fitz  # PyMuPDF
import docx
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

from tests.fixtures.sample_data import (
    VALID_AADHAAR_CATALOG,
    VALID_PAN_CATALOG,
    VALID_CARD_CATALOG,
    VALID_PHONE_CATALOG,
    VALID_EMAIL_CATALOG,
    VALID_NAMES_CATALOG,
    VALID_ADDRESSES_CATALOG
)


def create_txt(content: str) -> bytes:
    """Create UTF-8 encoded text file bytes."""
    return content.encode("utf-8")


def create_csv(rows: List[List[Any]]) -> bytes:
    """Create RFC 4180 compliant CSV file bytes."""
    output = io.StringIO()
    writer = csv.writer(output)
    for row in rows:
        writer.writerow(row)
    return output.getvalue().encode("utf-8")


def create_xlsx(sheets_data: Dict[str, List[List[Any]]]) -> bytes:
    """Create valid multi-sheet Excel spreadsheet bytes."""
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    for sheet_name, rows in sheets_data.items():
        ws = wb.create_sheet(title=sheet_name)
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                # Apply header styling for first row
                if r_idx == 1:
                    cell.font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def create_docx(
    paragraphs: List[str],
    tables_data: Optional[List[List[List[str]]]] = None,
    headings: Optional[List[Tuple[str, int]]] = None
) -> bytes:
    """Create valid Word document (.docx) bytes with paragraphs, runs, and optional tables."""
    doc = docx.Document()

    if headings:
        for title, level in headings:
            doc.add_heading(title, level=level)

    for p_text in paragraphs:
        doc.add_paragraph(p_text)

    if tables_data:
        for table_matrix in tables_data:
            if not table_matrix:
                continue
            rows = len(table_matrix)
            cols = len(table_matrix[0]) if rows > 0 else 0
            tbl = doc.add_table(rows=rows, cols=cols)
            tbl.style = "Table Grid"
            for r_idx, row in enumerate(table_matrix):
                for c_idx, val in enumerate(row):
                    tbl.cell(r_idx, c_idx).text = str(val)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_pdf(pages_text: List[List[Tuple[float, float, str]]]) -> bytes:
    """Create valid text-based PDF document bytes using PyMuPDF.

    Args:
        pages_text: A list of pages, where each page is a list of (x, y, text) tuples.
    """
    doc = fitz.open()
    for page_items in pages_text:
        # Standard Letter size: 612 x 792 pt
        page = doc.new_page(width=612, height=792)
        for x, y, text in page_items:
            page.insert_text(fitz.Point(x, y), text, fontname="helv", fontsize=11)

    buf = doc.tobytes()
    doc.close()
    return buf


def generate_simple_pdf(text: str) -> bytes:
    """Create a single-page PDF containing the provided text."""
    lines = text.split("\n")
    items = []
    y = 72.0
    for line in lines:
        if line.strip():
            items.append((72.0, y, line))
        y += 18.0
    return create_pdf([items])


def generate_simple_docx(text: str) -> bytes:
    """Create a single-paragraph or multi-paragraph DOCX from text."""
    paragraphs = [p for p in text.split("\n") if p.strip()]
    return create_docx(paragraphs)


def generate_simple_csv(header: List[str], data_rows: List[List[Any]]) -> bytes:
    """Create standard CSV with headers and data."""
    return create_csv([header] + data_rows)


def generate_simple_xlsx(sheet_name: str, header: List[str], data_rows: List[List[Any]]) -> bytes:
    """Create single-sheet XLSX with headers and data."""
    return create_xlsx({sheet_name: [header] + data_rows})


# =====================================================================
# TIER 4 REAL-WORLD SYNTHETIC DOCUMENT GENERATORS
# =====================================================================

def generate_synthetic_resume_pdf() -> bytes:
    """Synthesize realistic 2-page candidate resume with planted PII.

    Planted PII:
    - Person: Aarav Sharma
    - Email: aarav.sharma@domain.com
    - Phone: +91 98765 43210
    - Aadhaar: 3675 9832 4511
    - PAN: ABCPE1234F
    - Address: Flat 402, Lotus Towers, MG Road, Koramangala, Bengaluru, Karnataka 560034
    """
    doc = fitz.open()

    # Page 1: Header, Contact, Summary, Education, Experience
    page1 = doc.new_page(width=612, height=792)
    p1_items = [
        (72.0, 60.0, "Aarav Sharma - Senior Software Engineer"),
        (72.0, 80.0, "Contact: +91 98765 43210 | Email: aarav.sharma@domain.com"),
        (72.0, 100.0, "Address: Flat 402, Lotus Towers, MG Road, Koramangala, Bengaluru, Karnataka 560034"),
        (72.0, 130.0, "PROFESSIONAL SUMMARY"),
        (72.0, 150.0, "Accomplished backend engineer with 6+ years specializing in distributed systems,"),
        (72.0, 168.0, "cloud architecture, and microservices. Proven track record in high-throughput APIs."),
        (72.0, 200.0, "EDUCATION & CREDENTIALS"),
        (72.0, 220.0, "B.Tech in Computer Science, Indian Institute of Technology (2018 - 2022)"),
        (72.0, 238.0, "Cumulative GPA: 3.85 / 4.0 | Dean's Honor List"),
        (72.0, 270.0, "TECHNICAL SKILLS"),
        (72.0, 290.0, "Languages: Python, Go, TypeScript, SQL, Bash"),
        (72.0, 308.0, "Frameworks: FastAPI, Presidio, PyTorch, React, Docker, Kubernetes"),
        (72.0, 340.0, "PROFESSIONAL EXPERIENCE"),
        (72.0, 360.0, "Staff Infrastructure Engineer | TechCorp Global (2022 - Present)"),
        (72.0, 378.0, "- Architected zero-trust identity verification microservice serving 10M daily requests."),
        (72.0, 396.0, "- Reduced document parsing latency by 45% using PyMuPDF streaming parsers."),
        (72.0, 426.0, "Software Engineer | DataScale Systems (2020 - 2022)"),
        (72.0, 444.0, "- Built distributed log indexing pipeline with Apache Kafka and Elasticsearch.")
    ]
    for x, y, text in p1_items:
        page1.insert_text(fitz.Point(x, y), text, fontname="helv", fontsize=10)

    # Page 2: Identity Verification Section (Planted Aadhaar & PAN) & References
    page2 = doc.new_page(width=612, height=792)
    p2_items = [
        (72.0, 60.0, "GOVERNMENT IDENTIFICATION & COMPLIANCE VERIFICATION"),
        (72.0, 85.0, "Candidate Aadhaar Card Number: 3675 9832 4511"),
        (72.0, 110.0, "Income Tax Permanent Account Number (PAN): ABCPE1234F"),
        (72.0, 135.0, "Status: Verified through UIDAI and NSDL tax portal integrations."),
        (72.0, 175.0, "PROJECT PORTFOLIO"),
        (72.0, 195.0, "1. SentinelDoc Data Privacy Engine: Real-time redaction for enterprise files."),
        (72.0, 215.0, "2. High-Throughput Luhn and Verhoeff Checksum Verification Engine."),
        (72.0, 255.0, "REFERENCES"),
        (72.0, 275.0, "Available upon formal corporate compliance review request.")
    ]
    for x, y, text in p2_items:
        page2.insert_text(fitz.Point(x, y), text, fontname="helv", fontsize=10)

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def generate_synthetic_invoice_xlsx() -> bytes:
    """Synthesize corporate tax invoice spreadsheet with planted PII and non-PII financial data.

    Planted PII:
    - Customer: Priya Ramanathan
    - Address: Plot 12, Sector 18, Electronic City Phase 1, Bangalore 560100
    - Card: 4532 0151 1283 0366
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tax_Invoice"

    # Header section
    ws["A1"] = "TAX INVOICE - SENTINEL TECH SOLUTIONS PVT LTD"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "GSTIN: 29ABCDE1234F1Z5 | HSN Code: 998313 | Invoice No: INV-2026-0905"

    # Customer & Payment Info (Planted PII)
    ws["A4"] = "Customer Name:"
    ws["B4"] = "Priya Ramanathan"
    ws["A5"] = "Billing Address:"
    ws["B5"] = "Plot 12, Sector 18, Electronic City Phase 1, Bangalore 560100"
    ws["A6"] = "Payment Method (Credit Card):"
    ws["B6"] = "4532 0151 1283 0366"
    ws["A7"] = "Payment Status:"
    ws["B7"] = "Paid via Visa Online Gateway"

    # Line Items Table
    headers = ["Item #", "Description", "HSN Code", "Qty", "Unit Price (INR)", "Tax Rate (%)", "Total (INR)"]
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=9, column=col_idx, value=h)
        c.font = Font(bold=True)

    items = [
        [1, "Enterprise Privacy Scanner License (Annual)", "998313", 1, 120000.0, 18.0, 141600.0],
        [2, "Custom Presidio Recognizer Module", "998314", 2, 35000.0, 18.0, 82600.0],
        [3, "Multi-Format Redaction Plugin (PDF/DOCX)", "998315", 1, 45000.0, 18.0, 53100.0],
        [4, "High-Availability API Gateway Setup", "998316", 1, 60000.0, 18.0, 70800.0],
        [5, "Security Audit and Compliance Sign-Off", "998317", 1, 80000.0, 18.0, 94400.0],
        [6, "Onsite Engineering Training Session", "998318", 3, 15000.0, 18.0, 53100.0],
        [7, "24/7 SLA Support Tier 1", "998319", 12, 10000.0, 18.0, 141600.0],
        [8, "Data Loss Prevention Verification Kit", "998320", 1, 25000.0, 18.0, 29500.0]
    ]

    for r_offset, row in enumerate(items, start=10):
        for c_offset, val in enumerate(row, start=1):
            ws.cell(row=r_offset, column=c_offset, value=val)

    # Subtotal and Grand Total
    ws["F19"] = "Subtotal:"
    ws["G19"] = 565000.0
    ws["F20"] = "GST (18%):"
    ws["G20"] = 101700.0
    ws["F21"] = "Grand Total:"
    ws["G21"] = 666700.0
    ws["F21"].font = Font(bold=True)
    ws["G21"].font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_synthetic_customer_database_csv(row_count: int = 50) -> bytes:
    """Synthesize 50-row customer database CSV containing planted PII alongside tricky non-PII values."""
    header = [
        "cust_id",
        "full_name",
        "email",
        "phone",
        "aadhaar_no",
        "pan_no",
        "card_no",
        "billing_address",
        "order_id",
        "zip_code"
    ]

    rows = []
    # Seed known planted rows
    rows.append([
        "CUST-001",
        "Aarav Sharma",
        "aarav.sharma@domain.com",
        "+91 98765 43210",
        "3675 9832 4511",
        "ABCPE1234F",
        "4532 0151 1283 0366",
        "Flat 402, Lotus Towers, MG Road, Bengaluru 560034",
        "1002938475",  # non-PII order id
        "560034"        # non-PII postal zip
    ])
    rows.append([
        "CUST-002",
        "Priya Ramanathan",
        "compliance_officer+audit@enterprise.org.in",
        "+91-9123456789",
        "982341567897",
        "XYZCA5678B",
        "5424-1801-2345-6789",
        "Plot 12, Sector 18, Electronic City Phase 1, Bangalore 560100",
        "ORD-123456789012",  # non-PII order id
        "560100"
    ])
    rows.append([
        "CUST-003",
        "Vikramaditya Mukherjee",
        "harsh.patel@sentineldoc.com",
        "9876543210",
        "4321-8765-2109",
        "BNZPK9876M",
        "378282246310005",
        "24B Park Street, Near City Center, Kolkata 700016",
        "1725495000",   # non-PII epoch timestamp
        "700016"
    ])

    # Fill remaining rows up to row_count with non-PII business accounts
    for i in range(4, row_count + 1):
        rows.append([
            f"CUST-{i:03d}",
            f"Department Account {i}",
            f"dept_{i}@sentineldoc.com",
            f"080-234567{i:02d}",
            "3675 9832 4512" if i % 2 == 0 else "1234 5678 9012",  # Invalid Aadhaar (Verhoeff fails / starts with 1)
            f"ABCD{i:05d}E",  # Invalid PAN pattern (4 letters + 5 digits)
            "4532 0151 1283 0367",  # Invalid card (Luhn fails)
            f"Suite {i*10}, Corporate Plaza, Industrial Area",
            f"ORD-{1000000000 + i}",
            f"560{i:03d}"
        ])

    return create_csv([header] + rows)


def generate_synthetic_confidential_memo_docx() -> bytes:
    """Synthesize confidential executive severance agreement memo in DOCX.

    Planted PII:
    - Person: Vikramaditya Mukherjee
    - PAN: BNZPK9876M
    - Address: 24B Park Street, Near City Center, Kolkata, West Bengal 700016
    - Email: harsh.patel@sentineldoc.com
    """
    headings = [
        ("STRICTLY CONFIDENTIAL — EXECUTIVE SEVERANCE AGREEMENT", 1),
        ("Reference: HR/EXEC/2026/0905", 2)
    ]

    paragraphs = [
        "This Mutual Separation and Severance Agreement ('Agreement') is entered into between Sentinel Technologies Inc. and the departing executive.",
        "Employee Party: Vikramaditya Mukherjee, Chief Information Security Officer.",
        "Residential Address on Record: 24B Park Street, Near City Center, Kolkata, West Bengal 700016.",
        "Income Tax Identification (PAN): BNZPK9876M.",
        "Official Notification Email: harsh.patel@sentineldoc.com.",
        "1. Severance Consideration: The company shall disburse the gross sum of INR 45,00,000 via wire transfer.",
        "2. Non-Disclosure & Confidentiality: The executive agrees not to disclose proprietary vulnerability research or scanning algorithms.",
        "3. Surrender of Company Assets: All encrypted laptop keys, smart cards, and building access fobs must be returned immediately."
    ]

    table_data = [
        [
            ["Section", "Obligation", "Sign-Off Status"],
            ["IP Assignment", "All custom recognizers assigned to SentinelDoc", "Completed"],
            ["Tax Clearance", "TDS deducted under PAN BNZPK9876M", "Verified"],
            ["Final Settlement", "Wiring to designated account", "Pending Authorization"]
        ]
    ]

    return create_docx(paragraphs, tables_data=table_data, headings=headings)


def generate_synthetic_server_audit_txt(lines_count: int = 500) -> bytes:
    """Synthesize 500-line server audit log file containing planted customer leaks and non-PII telemetry."""
    lines = []
    lines.append("# SentinelDoc Production API Gateway Access & Error Audit Log")
    lines.append("# Format: TIMESTAMP [LEVEL] [TRACE_ID] MESSAGE")

    for i in range(1, lines_count + 1):
        ts = f"2026-09-05T02:{i % 60:02d}:{(i * 3) % 60:02d}.{i % 1000:03d}Z"
        trace_id = f"req-{i:06d}"

        if i == 50:
            lines.append(f"{ts} [WARN] [{trace_id}] Customer data leaked in query params: user=aarav.sharma@domain.com")
        elif i == 150:
            lines.append(f"{ts} [ERROR] [{trace_id}] Payment gateway error for card 4532 0151 1283 0366 - connection timeout")
        elif i == 250:
            lines.append(f"{ts} [INFO] [{trace_id}] Identity verification attempt with Aadhaar 3675 9832 4511 from IP 192.168.1.100")
        elif i == 350:
            lines.append(f"{ts} [WARN] [{trace_id}] Inbound mobile contact request logged: phone=+91 98765 43210")
        elif i == 450:
            lines.append(f"{ts} [INFO] [{trace_id}] Tax documentation submitted for PAN ABCPE1234F")
        else:
            # Routine non-PII operational log lines
            lines.append(
                f"{ts} [INFO] [{trace_id}] HTTP 200 GET /api/v1/health status=ok latency={12 + (i % 15)}ms memory_mb=256"
            )

    return create_txt("\n".join(lines) + "\n")
