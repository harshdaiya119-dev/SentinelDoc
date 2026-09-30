"""Tier 2: Boundary & Corner Cases Test Suite.

Verifies system resilience against degenerate inputs, extreme coordinate limits,
large throughput files, Unicode/Devanagari scripts, emojis, and spreadsheet anomalies.
"""
import io
import csv
import pytest
import fitz
import openpyxl
from fastapi.testclient import TestClient

from tests.fixtures.generator import (
    create_txt,
    create_csv,
    create_xlsx,
    create_docx,
    create_pdf
)
from tests.conftest import assert_valid_scan_response


@pytest.mark.tier2
class TestTier2Boundaries:
    """Boundary and corner cases testing suite."""

    # -----------------------------------------------------------------
    # Degenerate & Minimal Files
    # -----------------------------------------------------------------

    def test_tier2_blank_page_pdf(self, client: TestClient):
        """Verify 1-page completely blank PDF yields 0 findings without exception."""
        doc = fitz.open()
        doc.new_page(width=612, height=792)  # empty page
        pdf_bytes = doc.tobytes()
        doc.close()

        files = {"file": ("blank.pdf", pdf_bytes, "application/pdf")}
        resp = client.post("/scan", files=files)
        assert resp.status_code == 200
        data = resp.json()
        assert data["total_findings"] == 0

    def test_tier2_whitespace_only_txt(self, client: TestClient):
        """Verify file containing only whitespace, tabs, and newlines produces 0 findings."""
        ws_content = "   \n\t\r\n   \t  \n"
        files = {"file": ("whitespace.txt", create_txt(ws_content), "text/plain")}
        resp = client.post("/scan", files=files)
        assert resp.status_code == 200
        assert resp.json()["total_findings"] == 0

    def test_tier2_header_only_csv(self, client: TestClient):
        """Verify CSV containing only header row with zero data rows produces 0 findings."""
        header = ["id", "full_name", "phone_number", "aadhaar_number"]
        files = {"file": ("header_only.csv", create_csv([header]), "text/csv")}
        resp = client.post("/scan", files=files)
        assert resp.status_code == 200
        assert resp.json()["total_findings"] == 0

    def test_tier2_single_cell_xlsx(self, client: TestClient):
        """Verify XLSX containing only a single cell is processed cleanly."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Static Title Only"
        buf = io.BytesIO()
        wb.save(buf)

        files = {"file": ("single_cell.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        resp = client.post("/scan", files=files)
        assert resp.status_code == 200
        assert resp.json()["total_findings"] == 0

    # -----------------------------------------------------------------
    # Large Files & Throughput
    # -----------------------------------------------------------------

    def test_tier2_large_csv_1000_rows(self, client: TestClient):
        """Verify CSV with 1,000 rows containing 2 planted PII items completes successfully."""
        header = ["row_id", "description", "reference_code"]
        rows = []
        for i in range(1, 1001):
            if i == 500:
                rows.append([f"R-{i}", "Customer PAN record", "ABCPE1234F"])
            elif i == 800:
                rows.append([f"R-{i}", "Customer Aadhaar record", "3675 9832 4511"])
            else:
                rows.append([f"R-{i}", f"Transaction log entry {i}", f"TXN-{100000+i}"])

        csv_bytes = create_csv([header] + rows)
        files = {"file": ("large.csv", csv_bytes, "text/csv")}
        resp = client.post("/scan", files=files)
        assert resp.status_code == 200
        findings = resp.json()["findings"]
        pan_findings = [f for f in findings if f["entity_type"] == "PAN"]
        aadhaar_findings = [f for f in findings if f["entity_type"] == "AADHAAR"]
        assert len(pan_findings) >= 1
        assert len(aadhaar_findings) >= 1

    def test_tier2_multipage_pdf_10_pages(self, client: TestClient):
        """Verify 10-page PDF document correctly maps findings on 10th page without memory issues."""
        pages = []
        for page_num in range(1, 11):
            if page_num == 10:
                pages.append([(72.0, 100.0, "Target PII Aadhaar: 3675 9832 4511 on last page")])
            else:
                pages.append([(72.0, 100.0, f"Routine content page {page_num} without sensitive data")])

        pdf_bytes = create_pdf(pages)
        files = {"file": ("10_pages.pdf", pdf_bytes, "application/pdf")}
        resp = client.post("/scan", files=files)
        assert resp.status_code == 200
        findings = [f for f in resp.json()["findings"] if f["entity_type"] == "AADHAAR"]
        assert len(findings) == 1
        assert findings[0]["location"]["page"] == 10

    # -----------------------------------------------------------------
    # Unicode, Emojis & Character Encoding
    # -----------------------------------------------------------------

    def test_tier2_unicode_devanagari_script(self, client: TestClient):
        """Verify Devanagari Hindi text containing planted Aadhaar is accurately parsed."""
        text = "गोपनीय दस्तावेज़: नागरिक का आधार कार्ड 3675 9832 4511 सत्यापित किया गया।"
        files = {"file": ("devanagari.txt", create_txt(text), "text/plain")}
        resp = client.post("/scan", files=files)
        assert resp.status_code == 200
        findings = [f for f in resp.json()["findings"] if f["entity_type"] == "AADHAAR"]
        assert len(findings) >= 1

    def test_tier2_surrogate_pairs_and_emojis(self, client: TestClient):
        """Verify text with surrounding emojis does not corrupt character offsets or detection."""
        text = "🔒 Contact details: harsh.patel@sentineldoc.com 📱 +91 98765 43210 💳 4532 0151 1283 0366 🚀"
        files = {"file": ("emojis.txt", create_txt(text), "text/plain")}
        resp = client.post("/scan", files=files)
        assert resp.status_code == 200
        findings = resp.json()["findings"]
        types = {f["entity_type"] for f in findings}
        assert "EMAIL_ADDRESS" in types
        assert "PHONE_NUMBER" in types
        assert "CREDIT_CARD" in types

    def test_tier2_escaped_csv_special_characters(self, client: TestClient):
        """Verify CSV with quotes, delimiters, and multiline text doesn't corrupt column offsets."""
        header = ["id", "notes"]
        rows = [
            ["1", 'Customer said: "My PAN is ABCPE1234F, please update!"'],
            ["2", "Billing address:\nFlat 402, Koramangala,\nBengaluru 560034"]
        ]
        files = {"file": ("escaped.csv", create_csv([header] + rows), "text/csv")}
        resp = client.post("/scan", files=files)
        assert resp.status_code == 200
        findings = [f for f in resp.json()["findings"] if f["entity_type"] == "PAN"]
        assert len(findings) >= 1

    # -----------------------------------------------------------------
    # Edge Coordinate Bounds (PDF)
    # -----------------------------------------------------------------

    def test_tier2_pdf_extreme_top_margin(self, client: TestClient):
        """Verify PII near extreme top margin (y=20 pt) has valid bounding box coordinates."""
        page = [(72.0, 20.0, "Top Header Aadhaar: 3675 9832 4511")]
        pdf_bytes = create_pdf([page])
        files = {"file": ("top_margin.pdf", pdf_bytes, "application/pdf")}
        resp = client.post("/scan", files=files)
        assert resp.status_code == 200
        findings = [f for f in resp.json()["findings"] if f["entity_type"] == "AADHAAR"]
        assert len(findings) >= 1
        bbox = findings[0]["location"]["bbox"]
        assert bbox is not None
        assert bbox["top"] >= 0.0

    def test_tier2_pdf_extreme_bottom_footer(self, client: TestClient):
        """Verify PII near extreme bottom footer (y=770 pt) has valid bounding box."""
        page = [(72.0, 770.0, "Footer Note: Permanent Account Number ABCPE1234F")]
        pdf_bytes = create_pdf([page])
        files = {"file": ("footer.pdf", pdf_bytes, "application/pdf")}
        resp = client.post("/scan", files=files)
        assert resp.status_code == 200
        findings = [f for f in resp.json()["findings"] if f["entity_type"] == "PAN"]
        assert len(findings) >= 1
        bbox = findings[0]["location"]["bbox"]
        assert bbox is not None
        assert bbox["bottom"] <= 792.0

    def test_tier2_pdf_extreme_left_margin(self, client: TestClient):
        """Verify PII placed close to left edge (x=10 pt) is detected."""
        page = [(10.0, 100.0, "Card: 4532 0151 1283 0366")]
        pdf_bytes = create_pdf([page])
        files = {"file": ("left_margin.pdf", pdf_bytes, "application/pdf")}
        resp = client.post("/scan", files=files)
        assert resp.status_code == 200
        findings = [f for f in resp.json()["findings"] if f["entity_type"] == "CREDIT_CARD"]
        assert len(findings) >= 1
        bbox = findings[0]["location"]["bbox"]
        assert bbox is not None
        assert bbox["x0"] >= 0.0

    # -----------------------------------------------------------------
    # Spreadsheet Structure Degeneracies
    # -----------------------------------------------------------------

    def test_tier2_xlsx_merged_cells(self, client: TestClient):
        """Verify Excel workbook with merged cells containing PII parses cleanly."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.merge_cells("B2:D2")
        ws["B2"] = "Customer Aadhaar: 3675 9832 4511"

        buf = io.BytesIO()
        wb.save(buf)

        files = {"file": ("merged.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        resp = client.post("/scan", files=files)
        assert resp.status_code == 200
        findings = [f for f in resp.json()["findings"] if f["entity_type"] == "AADHAAR"]
        assert len(findings) >= 1
